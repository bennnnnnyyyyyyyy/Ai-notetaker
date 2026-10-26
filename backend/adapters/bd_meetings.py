"""Read-only adapter for the local BD Meetings workbook."""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from backend.config import settings

_HYPERLINK_RE = re.compile(r'HYPERLINK\("([^"]+)"', re.IGNORECASE)


class BDMeetingsSetupError(RuntimeError):
    pass


def _workbook_path() -> Path:
    return Path(settings.bd_meetings_workbook_path)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return value.isoformat() if isinstance(value, datetime) else str(value).strip()


def _links(value: Any) -> tuple[str, str]:
    text = _cell_text(value)
    match = _HYPERLINK_RE.search(text)
    url = match.group(1) if match else (text if text.startswith("http") else "")
    return (
        url if "meet.google.com" in url else "",
        url if "calendar.google.com" in url else "",
    )


def _load_rows() -> list[dict[str, Any]]:
    path = _workbook_path()
    if not path.is_file():
        raise BDMeetingsSetupError("BD Meetings workbook was not found on this PC.")
    try:
        import openpyxl
    except ImportError as exc:
        raise BDMeetingsSetupError("Install openpyxl to read the BD Meetings workbook.") from exc

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet_name = settings.bd_meetings_worksheet
    if sheet_name not in workbook.sheetnames:
        raise BDMeetingsSetupError(f"The workbook has no '{sheet_name}' worksheet.")
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [_cell_text(value).lower() for value in next(rows, ())]
    positions = {header: index for index, header in enumerate(headers) if header}

    def value(row: tuple, header: str) -> Any:
        index = positions.get(header)
        return row[index] if index is not None and index < len(row) else None

    result = []
    for row_number, row in enumerate(rows, start=2):
        company = _cell_text(value(row, "company name"))
        contact = _cell_text(value(row, "authorized person"))
        email = _cell_text(value(row, "email"))
        if not any((company, contact, email)):
            continue
        meeting_value = value(row, "meeting time")
        meet_url, calendar_url = _links(meeting_value)
        result.append({
            "source_row": row_number,
            "campaign_label": _cell_text(value(row, "sub")),
            "opener_name": _cell_text(value(row, "opener")),
            "source_status": _cell_text(value(row, "status")),
            "date_added": _cell_text(value(row, "date added")),
            "company_name": company,
            "client_name": contact,
            "email": email,
            "meeting_time": _cell_text(meeting_value),
            "opener_summary": _cell_text(value(row, "opener summary")),
            "closer_notes": _cell_text(value(row, "closer's notes")),
            "calendar_url": calendar_url,
            "meet_url": meet_url,
            "ready_to_import": bool(meet_url),
        })
    return result


def get_status() -> dict[str, Any]:
    path = _workbook_path()
    ready = path.is_file()
    return {
        "status": "ready" if ready else "needs_setup",
        "workbook_name": path.name,
        "worksheet": settings.bd_meetings_worksheet,
        "detail": "Read-only BD Meetings workbook is available." if ready else "Configure BD_MEETINGS_WORKBOOK_PATH.",
    }


def list_bookings(limit: int = 200) -> list[dict[str, Any]]:
    return _load_rows()[: max(1, min(limit, 500))]
